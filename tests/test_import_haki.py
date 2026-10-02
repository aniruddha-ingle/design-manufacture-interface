"""The Haki deck importer, on synthetic decks shaped like Haki's (never the real ones)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Cm

from dmi.importers.haki_deck import import_brief, import_revisions, read_slides
from dmi.spec import dump

PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6300010000050001"
    "0d0a2db40000000049454e44ae426082"
)


def deck(tmp: Path, name: str, slides: list[dict]) -> Path:
    prs = Presentation()
    img = tmp / "px.png"
    img.write_bytes(PNG)
    for i, s in enumerate(slides):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        sl = prs.slides.add_slide(layout)
        sl.shapes.title.text = s["title"]
        if i > 0 and "body" not in s:
            sl.placeholders[1].element.getparent().remove(sl.placeholders[1].element)
        if "body" in s:
            sl.placeholders[1].text = s["body"]
        if i == 0:
            sl.placeholders[1].element.getparent().remove(sl.placeholders[1].element)
        for k, text in enumerate(s.get("texts", [])):
            tb = sl.shapes.add_textbox(Cm(15), Cm(3 + 2 * k), Cm(8), Cm(1.5))
            tb.text_frame.text = text
        for k in range(s.get("pictures", 0)):
            sl.shapes.add_picture(str(img), Cm(1 + 6 * k), Cm(4), Cm(5), Cm(5))
    out = tmp / name
    prs.save(out)
    return out


@pytest.fixture
def brief(tmp_path: Path) -> Path:
    return deck(
        tmp_path,
        "brief.pptx",
        [
            {"title": "“Example” Cap - Olive"},
            {"title": "Shape References", "pictures": 2},
            {"title": "Front Embroidery - Puff", "pictures": 1},
            {
                "title": "Side Embroidery (Cream + Rust)",
                "pictures": 1,
                "texts": ["Shape is not EXACT in the mockup, use Shape References"],
            },
            {"title": "Back of Cap (Label)", "pictures": 2},
            {
                "title": "Construction Details",
                "body": "STRUCTURED, LOW CROWN, LONG BRIM FIT\nCOTTON CANVAS FABRIC\n"
                "BROWN SUEDE STRAP & BRASS HARDWARE\nEMBROIDERY ON FRONT AND SIDE",
            },
        ],
    )


@pytest.fixture
def revisions(tmp_path: Path) -> Path:
    return deck(
        tmp_path,
        "rev.pptx",
        [
            {"title": "“Example” Cap - Olive", "texts": ["Revisions - Sample 3"], "pictures": 1},
            {
                "title": "#1: Change position + size of front embroidery",
                "pictures": 2,
                "texts": ["Before:", "After:", "Move the mark left. Make it 5% bigger."],
            },
            {
                "title": "#2: Change crown height",
                "pictures": 1,
                "texts": ["4.2 cm", "Current sample is 4.6cm, adjust to 4.2cm"],
            },
            {
                "title": "#3. Back of Cap (Label)",
                "pictures": 1,
                "texts": ["Please use a smooth buckle"],
            },
            {
                "title": "#3. Back of Cap (Label)",
                "pictures": 1,
                "texts": ["The buckle scratches the strap"],
            },
            {
                "title": "#4. Adjust brim length",
                "pictures": 1,
                "texts": ["5.7 cm", "Current sample brim is 6.1 cm, adjust down to 5.7 cm"],
            },
            {"title": "#5. Add Inner Label", "pictures": 3},
        ],
    )


def test_read_slides_hashes_not_names(brief: Path):
    deck_id, slides = read_slides(brief)
    assert deck_id.startswith("deck-") and "example" not in deck_id.lower()
    assert [s.pictures for s in slides] == [0, 2, 1, 1, 2, 0]


def test_brief_deck_becomes_spec_v1(brief: Path):
    p = import_brief(brief, "pre:haki:cap-01")
    assert (p.spec_version, p.colourway, p.category) == (1, "olive", "headwear")
    front, side = p.placements
    assert (front.id, front.technique, front.location) == (
        "front-embroidery",
        "puff-embroidery",
        "front",
    )
    assert (side.location, side.technique) == ("side", "flat-embroidery")
    assert [e.colour.name for e in side.elements] == ["Cream", "Rust"]
    assert front.width.mm is None and front.position is None  # gaps, not guesses
    kinds = [r.kind for r in p.references]
    assert kinds.count("shape") == 2 and kinds.count("mockup") == 2 and kinds.count("detail") == 2
    side_mock = next(r for r in p.references if r.caption == "side-embroidery mockup")
    assert side_mock.authority == "approximate"
    c = p.construction
    assert (c.structure, c.crown) == ("structured", "low")
    assert len(c.notes) == 4  # every line kept verbatim
    assert {m.id for m in p.materials} == {"shell", "strap"}
    assert p.hardware[0].kind == "clasp"
    assert p.labels[0].kind == "back-of-cap"


def test_revision_deck_becomes_a_sample_round(brief: Path, revisions: Path):
    p = import_revisions(revisions, import_brief(brief, "pre:haki:cap-01"))
    assert p.spec_version == 2
    (r,) = p.sample_rounds
    assert (r.n, r.made_from_spec_version, r.applied_in_spec_version) == (3, 1, 2)
    by_n = {c.n: c for c in r.changes}
    assert sorted(by_n) == [1, 2, 3, 4, 5]
    assert (by_n[1].area, by_n[1].field) == ("placement-position", "placements.front-embroidery")
    assert (by_n[2].field, by_n[2].current_value.mm, by_n[2].target_value.mm) == (
        "poms.crown-height",
        46,
        42,
    )
    assert (by_n[4].field, by_n[4].target_value.mm) == ("poms.brim-length", 57)
    assert p.pom("brim-length").values["OS"].mm == 57
    assert (by_n[3].area, by_n[3].field) == ("hardware", "hardware.strap-hardware")
    assert len(by_n[3].image_refs) == 2  # the change spans two slides
    assert p.hardware[0].requirements == [
        "Please use a smooth buckle",
        "The buckle scratches the strap",
    ]
    assert (by_n[5].area, by_n[5].field) == ("label", "labels.inner-label")
    assert len(by_n[5].image_refs) == 3


def test_import_is_deterministic(brief: Path, revisions: Path):
    a = dump(import_revisions(revisions, import_brief(brief, "pre:haki:cap-01")))
    b = dump(import_revisions(revisions, import_brief(brief, "pre:haki:cap-01")))
    assert a == b


def test_not_a_revision_deck(brief: Path):
    with pytest.raises(ValueError, match="Revisions - Sample"):
        import_revisions(brief, import_brief(brief, "pre:haki:cap-01"))
