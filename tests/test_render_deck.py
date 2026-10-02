from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation

from dmi.render.deck import fmt, render_brief, render_revisions
from dmi.spec import load
from dmi.spec.models import Length

GOLDEN = Path(__file__).resolve().parent / "golden/example-cap.json"


def texts(data: bytes) -> list[list[str]]:
    prs = Presentation(io.BytesIO(data))
    return [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in prs.slides]


def test_fmt():
    assert fmt(Length(mm=75, tol_mm=2)) == '7.5 cm (3") ±0.2'
    assert fmt(Length(min_mm=540, max_mm=600, tol_mm=5)) == "54.0–60.0 cm ±0.5"
    assert fmt(Length(mm=230, tol_mm=3, tol_minus_mm=1)).endswith("+0.3/−0.1")
    assert fmt(Length()) is None and fmt(None) is None


def test_brief_follows_haki_structure_and_carries_numbers():
    slides = texts(render_brief(load(GOLDEN)))
    titles = [s[0] for s in slides]
    assert titles[0].startswith("Six-panel cap")
    assert "Shape References" in titles and "Construction Details" in titles
    assert any(t.startswith("Front Puff Embroidery (Cream + Forest)") for t in titles)
    front = next(s for s in slides if s[0].startswith("Front Puff"))
    body = "\n".join(front)
    assert (
        "Width: 6.0 cm" in body and "Position: bottom edge" in body and "madeira SYN-1001" in body
    )
    assert "Not specified yet" not in body  # the golden cap is complete
    assert any("Brim length: 7.5 cm" in "\n".join(s) for s in slides)


def test_gaps_are_shown_on_the_slide():
    p = load(GOLDEN)
    p.placements[0].width = Length()
    p.placements[0].position = None
    front = next(s for s in texts(render_brief(p)) if s[0].startswith("Front Puff"))
    gap = next(t for t in front if t.startswith("Not specified yet"))
    assert "width" in gap and "measured position" in gap


def test_revision_deck():
    p = load(GOLDEN)
    slides = texts(render_revisions(p, p.sample_rounds[0]))
    assert "Revisions - Sample 1" in slides[0]
    first = "\n".join(slides[1])
    assert slides[1][0] == "#1: Brim too long"
    assert "Current sample: 8.0 cm" in first and "Adjust to: 7.5 cm" in first
    assert "Why: brim reads oversized" in first


def test_same_spec_same_bytes():
    p = load(GOLDEN)
    assert render_brief(p) == render_brief(p)
    assert render_revisions(p, p.sample_rounds[0]) == render_revisions(p, p.sample_rounds[0])
