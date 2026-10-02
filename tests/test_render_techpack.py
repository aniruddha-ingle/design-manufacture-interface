from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook

from dmi.checks import check
from dmi.render.techpack import length, render_bom, render_pdf, tolerance, view
from dmi.spec import load
from dmi.spec.models import Length

GOLDEN = Path(__file__).resolve().parent / "golden/example-cap.json"


def flat(v: dict) -> list[dict]:
    return [
        c
        for s in v["sections"]
        for b in s["blocks"]
        if b["kind"] == "table"
        for r in b["rows"]
        for c in r
    ]


def test_cells():
    assert length(Length(mm=75)) == {"text": '7.5 cm / 3"', "gap": False}
    assert length(Length())["gap"] and tolerance(Length(mm=1))["gap"]
    assert tolerance(Length(mm=1, tol_mm=2, tol_minus_mm=1))["text"] == "+0.2 / −0.1 cm"


def test_golden_view_has_no_gaps_and_every_section():
    p = load(GOLDEN)
    v = view(p, check(p))
    titles = [s["title"] for s in v["sections"]]
    for t in (
        "Changes in v2",
        "References",
        "Measurements",
        "Bill of materials",
        "Construction",
        "Packaging and quantities",
        "Open gaps",
    ):
        assert t in titles, t
    assert sum(t.startswith("Artwork:") for t in titles) == 2
    assert not any(c["gap"] for c in flat(v))


def test_gaps_are_marked():
    p = load(GOLDEN)
    p.placements[0].position = None
    p.poms[2].values["OS"] = Length()
    v = view(p, check(p, severity="warn"))
    assert sum(c["gap"] for c in flat(v)) >= 3
    gaps = next(s for s in v["sections"] if s["title"] == "Open gaps")
    assert len(gaps["blocks"][1]["rows"]) == len(check(p))


def test_pdf_is_deterministic():
    p = load(GOLDEN)
    a, b = render_pdf(p, check(p)), render_pdf(p, check(p))
    assert a[:5] == b"%PDF-" and a == b


def test_pdf_with_images(tmp_path: Path):
    import struct
    import zlib

    def chunk(t: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + t + data + struct.pack(">I", zlib.crc32(t + data))

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
    )
    png += chunk(b"IEND", b"")
    (tmp_path / "a.png").write_bytes(png)
    p = load(GOLDEN)
    pdf = render_pdf(p, [], resolve=lambda ref: tmp_path / "a.png", root=tmp_path)
    assert pdf[:5] == b"%PDF-"


def test_bom_workbook():
    p = load(GOLDEN)
    data = render_bom(p, check(p))
    assert data == render_bom(p, check(p))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Bill of materials", "Measurements", "Construction", "Open gaps"]
    values = [c.value for row in wb["Bill of materials"].iter_rows() for c in row]
    assert "SYN-TWILL-280" in values
