"""Render a spec as Haki's own deck format: the brief deck and a revision deck per sample round.

The factory keeps the structure it already knows (title, shape references, a slide per artwork
placement, the back, construction details; numbered changes after a sample), and every slide
now also carries the spec's numbers: sizes and positions with tolerances, colour references,
the measurement table. Gaps are shown on the slide in a highlighted box, never left out.

Images are refs (never files in git): each shows as a labelled frame, or as the image when a
``resolve(ref) -> Path | None`` is given. Output is deterministic: same spec, same bytes.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

from dmi.spec.models import Length, Placement, Product, SampleRound
from dmi.spec.units import inch_fraction, mm_to_cm

Resolver = Callable[[str], Path | None]
FIXED = datetime(2000, 1, 1)
GAP_FILL = RGBColor(0xFF, 0xF1, 0xB8)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66)
W, H = Cm(25.4), Cm(14.29)  # Haki's decks are 16:9 at this size


def fmt(v: Length | None) -> str | None:
    """'8.0 cm (3 1/8") ±0.2' style, or None for a gap."""
    if v is None or not v.known:
        return None
    if v.mm is not None:
        s = f'{mm_to_cm(v.mm)} cm ({inch_fraction(v.mm)}")'
    else:
        s = f"{mm_to_cm(v.min_mm)}–{mm_to_cm(v.max_mm)} cm"
    if v.tol_mm is not None:
        minus = v.tol_minus_mm if v.tol_minus_mm is not None else v.tol_mm
        s += (
            f" ±{mm_to_cm(v.tol_mm, 2)}"
            if minus == v.tol_mm
            else f" +{mm_to_cm(v.tol_mm, 2)}/−{mm_to_cm(minus, 2)}"
        )
    return s


class _Deck:
    def __init__(self, resolve: Resolver | None):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.resolve = resolve

    def slide(self, title: str):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5])  # title only
        s.shapes.title.text = title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Cm(0.9), Cm(0.6), Cm(23.6), Cm(1.6)
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
        return s

    def text(self, s, x, y, w, h, lines: list[str], size=12, gap=False, colour=INK) -> None:
        tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        tf = tb.text_frame
        tf.word_wrap = True
        if gap:
            tb.fill.solid()
            tb.fill.fore_color.rgb = GAP_FILL
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.runs[0].font.size = Pt(size)
            p.runs[0].font.color.rgb = colour

    def image(self, s, ref: str, x, y, w, h) -> None:
        path = self.resolve(ref) if self.resolve else None
        if path is not None:
            s.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))
            return
        box = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        box.line.color.rgb = MUTED
        box.text_frame.text = f"image: {ref}"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED

    def images(self, s, refs: list[str], x0=0.7, y=3.0, width=13.0, height=10.5) -> None:
        if not refs:
            return
        w = width / len(refs) - 0.3
        for k, ref in enumerate(refs):
            self.image(s, ref, x0 + k * (w + 0.3), y, w, min(height, w * 1.3))

    def gaps(self, s, missing: list[str], y=11.6) -> None:
        if missing:
            self.text(s, 14.2, y, 10.4, 2.2, ["Not specified yet:", *missing], size=10, gap=True)

    def save(self) -> bytes:
        cp = self.prs.core_properties
        cp.created = cp.modified = cp.last_printed = FIXED
        cp.author = cp.last_modified_by = "design-manufacture-interface"
        cp.revision = 1
        buf = io.BytesIO()
        self.prs.save(buf)
        return _fixed_zip(buf.getvalue())


def _fixed_zip(data: bytes) -> bytes:
    """Rewrite the package with fixed entry times and order, so bytes depend only on content."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():  # [Content_Types].xml stays first
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(name))
    return out.getvalue()


def _placement_title(pl: Placement) -> str:
    loc = pl.location.replace("-", " ").title()
    tech = (pl.technique or "artwork").replace("-", " ").title()
    colours = " + ".join(e.colour.name for e in pl.elements if e.colour)
    return f"{loc} {tech}" + (f" ({colours})" if colours else "")


def _placement_lines(pl: Placement) -> tuple[list[str], list[str]]:
    lines, missing = [], []
    for label, v in (("Width", pl.width), ("Height", pl.height)):
        s = fmt(v)
        (lines.append(f"{label}: {s}") if s else missing.append(label.lower()))
    if pl.position and pl.position.dx_mm is not None and pl.position.dy_mm is not None:
        to = (pl.position.to or "item").replace("-", " ")
        lines.append(
            f"Position: {to} at {mm_to_cm(pl.position.dx_mm)} cm across, "
            f"{mm_to_cm(pl.position.dy_mm)} cm up from {pl.position.anchor.replace('-', ' ')}"
        )
    else:
        missing.append("measured position")
    for e in pl.elements:
        ref = f"{e.colour.system or ''} {e.colour.code or ''}".strip() if e.colour else ""
        thread = f", {e.thread}" if e.thread else ""
        lines.append(f"• {e.name}: {e.colour.name if e.colour else '?'} {ref}{thread}".rstrip())
        if not (e.colour and e.colour.code):
            missing.append(f"colour reference for {e.name}")
    if not pl.elements:
        missing.append("thread colours")
    if pl.technique == "puff-embroidery" and pl.foam_height_mm is None:
        missing.append("puff foam height")
    if not pl.artwork_ref:
        missing.append("artwork file")
    return [*lines, *pl.notes], missing


def render_brief(p: Product, resolve: Resolver | None = None) -> bytes:
    d = _Deck(resolve)
    s = d.slide(p.name)
    d.text(
        s,
        0.9,
        4.0,
        23.6,
        3.0,
        [f"Spec v{p.spec_version} · {p.id} · colourway {p.colourway}"],
        14,
        colour=MUTED,
    )

    shapes = [r for r in p.references if r.kind == "shape"]
    for k in range(0, len(shapes), 2):
        s = d.slide("Shape References")
        d.images(s, [r.ref for r in shapes[k : k + 2]], width=23.6)

    for pl in p.placements:
        s = d.slide(_placement_title(pl))
        mocks = [
            r.ref
            for r in p.references
            if r.kind == "mockup" and (r.caption or "").startswith(pl.id)
        ]
        d.images(s, mocks or ([pl.artwork_ref] if pl.artwork_ref else []))
        lines, missing = _placement_lines(pl)
        d.text(s, 14.2, 3.0, 10.4, 8.4, lines, 12)
        d.gaps(s, missing)

    if p.labels or p.hardware:
        s = d.slide("Back of Hat (Label)" if p.category == "headwear" else "Labels and Hardware")
        d.images(s, [r.ref for r in p.references if r.kind == "detail"])
        lines, missing = [], []
        for h in p.hardware:
            size = fmt(h.size)
            lines.append(f"• {h.kind}: {h.description}" + (f", {size}" if size else ""))
            lines += [f"   must: {r}" for r in h.requirements]
            if h.material is None or h.finish is None:
                missing.append(f"{h.kind} material/finish")
        for lb in p.labels:
            lines.append(f"• {lb.kind} label" + (f": {lb.material}" if lb.material else ""))
            if lb.position is None:
                missing.append(f"{lb.kind} label position")
        d.text(s, 14.2, 3.0, 10.4, 8.4, lines, 12)
        d.gaps(s, missing)

    s = d.slide("Measurements")
    rows = [f"{pm.name}: {fmt(pm.values.get(size)) or '—'}" for pm in p.poms for size in p.sizes]
    d.text(s, 0.9, 3.0, 23.6, 10.5, rows or ["No measurements yet"], 13, gap=not rows)

    s = d.slide("Construction Details")
    c = p.construction
    lines = [x.upper() for x in c.notes] or [
        ", ".join(
            x for x in (c.structure, f"{c.crown} crown" if c.crown else None, c.brim) if x
        ).upper()
    ]
    lines += [
        m.description.upper() for m in p.materials if m.description.upper() not in " ".join(lines)
    ]
    lines += [
        f"{sm.location}: {sm.stitch_type or ''} {f'{sm.spi:g} SPI' if sm.spi else ''}".strip()
        for sm in c.seams
    ]
    d.text(s, 0.9, 3.0, 23.6, 10.5, lines, 16)
    return d.save()


def render_revisions(p: Product, round_: SampleRound, resolve: Resolver | None = None) -> bytes:
    d = _Deck(resolve)
    s = d.slide(p.name)
    d.text(s, 0.9, 4.0, 23.6, 2.0, [f"Revisions - Sample {round_.n}"], 20)
    d.text(
        s,
        0.9,
        6.0,
        23.6,
        1.5,
        [f"Sample made from spec v{round_.made_from_spec_version}"],
        11,
        colour=MUTED,
    )
    for c in round_.changes:
        s = d.slide(f"#{c.n}: {c.summary}")
        d.images(s, c.image_refs[:3])
        lines = []
        cur, tgt = fmt(c.current_value), fmt(c.target_value)
        if cur or tgt:
            lines.append(f"Current sample: {cur or c.current or '—'}")
            lines.append(f"Adjust to: {tgt or c.target or '—'}")
        elif c.target:
            lines.append(c.target)
        if c.reason:
            lines.append(f"Why: {c.reason}")
        d.text(s, 14.2, 3.0, 10.4, 8.4, lines or ["(no instruction text)"], 13)
    return d.save()
