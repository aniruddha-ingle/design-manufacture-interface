"""The complete tech pack: a spec -> a PDF (Typst), a BOM workbook (xlsx) and the gap report.

Python turns the spec into a plain *view* (sections of paragraphs, tables and images, every
missing value marked as a gap); ``template.typ`` only lays the view out. So the template
stays generic and the logic stays tested. Same spec, same bytes: Typst gets no date, the
workbook gets fixed properties and zip entry times.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

import typst
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from dmi.checks import Finding
from dmi.spec.categories import POMS_BY_CATEGORY
from dmi.spec.models import Length, Position, Product
from dmi.spec.units import inch_fraction, mm_to_cm

Resolver = Callable[[str], Path | None]
TEMPLATE = Path(__file__).with_name("template.typ")
GAP = {"text": "not specified", "gap": True}


def cell(value, optional: bool = False) -> dict:
    """A value, or a gap. ``optional`` fields the checks don't require show "—" instead."""
    if value in (None, "", []):
        return {"text": "—", "gap": False} if optional else GAP
    return {"text": str(value), "gap": False}


def length(v: Length | None) -> dict:
    if v is None or not v.known:
        return GAP
    if v.mm is not None:
        s = f'{mm_to_cm(v.mm)} cm / {inch_fraction(v.mm)}"'
    else:
        cm = f"{mm_to_cm(v.min_mm)}–{mm_to_cm(v.max_mm)} cm"
        s = f'{cm} / {inch_fraction(v.min_mm)}–{inch_fraction(v.max_mm)}"'
    return cell(s)


def tolerance(v: Length | None, optional: bool = False) -> dict:
    if v is None or v.tol_mm is None:
        return cell(None, optional)
    minus = v.tol_minus_mm if v.tol_minus_mm is not None else v.tol_mm
    if minus == v.tol_mm:
        return cell(f"±{mm_to_cm(v.tol_mm, 2)} cm")
    return cell(f"+{mm_to_cm(v.tol_mm, 2)} / −{mm_to_cm(minus, 2)} cm")


def position(pos: Position | None) -> dict:
    if pos is None or pos.dx_mm is None or pos.dy_mm is None or pos.to is None:
        return GAP if pos is None else cell(f"from {pos.anchor} (incomplete: {pos.note or ''})")
    s = (
        f"{pos.to.replace('-', ' ')} at {mm_to_cm(pos.dx_mm)} cm across, "
        f"{mm_to_cm(pos.dy_mm)} cm up from {pos.anchor.replace('-', ' ')}"
    )
    return cell(s + (f" ({pos.note})" if pos.note else ""))


def colour(c) -> dict:
    if c is None:
        return GAP
    ref = " ".join(x for x in (c.system, c.code) if x)
    return cell(f"{c.name} ({ref})") if c.code else {"text": f"{c.name}: no reference", "gap": True}


def table(columns: list[str], rows: list[list[dict]], widths: list[str] | None = None) -> dict:
    return {"kind": "table", "columns": columns, "rows": rows, "widths": widths}


def para(text: str) -> dict:
    return {"kind": "para", "text": text}


def images(refs: list[tuple[str, str]], resolve: Resolver | None, root: Path) -> dict:
    out = []
    for ref, caption in refs:
        path = resolve(ref) if resolve else None
        rel = "/" + str(path.resolve().relative_to(root.resolve())) if path else None
        out.append({"path": rel, "ref": ref, "caption": caption})
    return {"kind": "images", "items": out}


def view(
    p: Product, findings: list[Finding], resolve: Resolver | None = None, root: Path = Path("/")
) -> dict:
    blocking = sum(f.severity == "block" for f in findings)
    cover = {
        "title": p.name,
        "rows": [
            ["Product id", p.id],
            ["Client", p.client],
            ["Category", p.category],
            ["Colourway", p.colourway],
            ["Season", p.season or "—"],
            ["Spec version", f"v{p.spec_version}"],
            ["Sizes", ", ".join(p.sizes)],
            ["Open gaps", f"{len(findings)} ({blocking} blocking)"],
        ],
    }
    s: list[dict] = []

    applied = [r for r in p.sample_rounds if r.applied_in_spec_version == p.spec_version]
    if applied:
        rows = []
        for r in applied:
            for c in r.changes:
                cur = length(c.current_value) if c.current_value else cell(c.current)
                tgt = length(c.target_value) if c.target_value else cell(c.target)
                rows.append(
                    [
                        cell(f"S{r.n} #{c.n}"),
                        cell(c.summary),
                        cell(c.field),
                        cur,
                        tgt,
                        cell(c.reason),
                    ]
                )
        s.append(
            {
                "title": f"Changes in v{p.spec_version}",
                "blocks": [
                    para("Applied from these sample rounds; the factory should check them first."),
                    table(
                        ["Change", "What", "Field", "Sample had", "Make it", "Why"],
                        rows,
                        ["auto", "2fr", "1.2fr", "1fr", "1.4fr", "1.4fr"],
                    ),
                ],
            }
        )

    if p.references:
        s.append(
            {
                "title": "References",
                "blocks": [
                    images(
                        [
                            (
                                r.ref,
                                f"{r.kind}{' · ' + r.caption if r.caption else ''} · "
                                + (r.authority or "exact or approximate? not stated"),
                            )
                            for r in p.references
                        ],
                        resolve,
                        root,
                    ),
                ],
            }
        )

    rows = []
    defs = {d.code: d for d in POMS_BY_CATEGORY.get(p.category, ())}
    for pm in p.poms:
        d = defs.get(pm.code)
        how = pm.how_to_measure or (d.how_to_measure if d else None)
        for size in p.sizes:
            v = pm.values.get(size)
            tol = tolerance(v, optional=bool(d and not d.needs_tolerance))
            rows.append([cell(pm.name), cell(size), length(v), tol, cell(how, optional=True)])
    s.append(
        {
            "title": "Measurements",
            "blocks": [
                table(
                    ["Point of measure", "Size", "Value", "Tolerance", "How to measure"],
                    rows or [[GAP] * 5],
                    ["1.3fr", "auto", "1.4fr", "1fr", "2fr"],
                )
            ],
        }
    )

    for pl in p.placements:
        mocks = [
            (r.ref, r.caption or "")
            for r in p.references
            if r.kind == "mockup" and (r.caption or "").startswith(pl.id)
        ]
        rows = [
            [cell("Technique"), cell(pl.technique)],
            [cell("Artwork file"), cell(pl.artwork_ref)],
            [cell("Width"), length(pl.width)],
            [cell("Width tolerance"), tolerance(pl.width)],
            [cell("Height"), length(pl.height)],
            [cell("Height tolerance"), tolerance(pl.height)],
            [cell("Position"), position(pl.position)],
            [cell("Stitch count"), cell(pl.stitch_count, optional=True)]
            if pl.technique and "embroider" in pl.technique or pl.technique == "chain-stitch"
            else None,
            [cell("Density"), cell(pl.density, optional=True)]
            if pl.technique and "embroider" in pl.technique
            else None,
            [cell("Foam height"), cell(f"{pl.foam_height_mm} mm" if pl.foam_height_mm else None)]
            if pl.technique == "puff-embroidery"
            else None,
        ]
        el = [
            [cell(e.name), colour(e.colour), cell(e.thread), cell(e.stitch_type)]
            for e in pl.elements
        ]
        blocks = [
            table(["", ""], [r for r in rows if r], ["1fr", "3fr"]),
            table(["Element", "Colour", "Thread", "Stitch"], el or [[GAP] * 4]),
        ]
        blocks += [para(n) for n in pl.notes]
        if mocks:
            blocks.append(images(mocks, resolve, root))
        s.append({"title": f"Artwork: {pl.location} · {pl.id}", "blocks": blocks})

    mats = [
        [
            cell(m.id),
            cell(m.role),
            cell(m.description),
            cell(", ".join(f"{f.percent:g}% {f.fibre}" for f in m.composition)),
            cell(f"{m.weight_gsm:g} gsm" if m.weight_gsm else None)
            if m.role == "shell"
            else length(m.thickness)
            if m.thickness
            else cell("—"),
            colour(m.colour),
            cell(m.supplier_ref),
        ]
        for m in p.materials
    ]
    trims = [
        [
            cell(t.id),
            cell(t.kind),
            cell(t.description),
            length(t.width),
            colour(t.colour) if not t.material_id else cell(f"as {t.material_id}"),
        ]
        for t in p.trims
    ]
    hw = [
        [
            cell(h.id),
            cell(h.kind),
            cell(h.description),
            cell(h.quantity),
            cell(h.material),
            cell(h.finish),
            length(h.size),
            cell(h.supplier_ref, optional=True),
            cell("; ".join(h.requirements) or "—"),
        ]
        for h in p.hardware
    ]
    lbs = [
        [
            cell(lb.id),
            cell(lb.kind),
            cell(lb.material),
            cell(lb.attachment),
            position(lb.position),
            length(lb.width),
            length(lb.height),
            cell(lb.content or lb.artwork_ref),
        ]
        for lb in p.labels
    ]
    s.append(
        {
            "title": "Bill of materials",
            "blocks": [
                para("Materials"),
                table(
                    [
                        "Id",
                        "Role",
                        "Description",
                        "Composition",
                        "Weight / thickness",
                        "Colour",
                        "Supplier",
                    ],
                    mats or [[GAP] * 7],
                ),
                para("Trims"),
                table(
                    ["Id", "Kind", "Description", "Width", "Colour"],
                    trims or [[cell("none")] + [cell("—")] * 4],
                ),
                para("Hardware"),
                table(
                    [
                        "Id",
                        "Kind",
                        "Description",
                        "Qty",
                        "Material",
                        "Finish",
                        "Size",
                        "Supplier",
                        "Requirements",
                    ],
                    hw or [[cell("none")] + [cell("—")] * 8],
                ),
                para("Labels"),
                table(
                    [
                        "Id",
                        "Kind",
                        "Material",
                        "Attachment",
                        "Position",
                        "Width",
                        "Height",
                        "Content / artwork",
                    ],
                    lbs or [[GAP] * 8],
                ),
            ],
        }
    )

    c = p.construction
    rows = [
        [cell(k.title()), cell(getattr(c, k))] for k in ("structure", "crown", "brim", "closure")
    ] + [[cell("Panels"), cell(c.panels)]]
    seams = [
        [
            cell(x.location),
            cell(x.stitch_type),
            cell(f"{x.spi:g}" if x.spi else None),
            cell(x.thread, optional=True),
            cell(x.rows, optional=True),
            cell(x.finish, optional=True),
        ]
        for x in c.seams
    ]
    s.append(
        {
            "title": "Construction",
            "blocks": [
                table(["", ""], rows, ["1fr", "3fr"]),
                table(["Seam", "Stitch", "SPI", "Thread", "Rows", "Finish"], seams or [[GAP] * 6]),
                *[para(n) for n in c.notes],
            ],
        }
    )

    pk = p.packaging
    s.append(
        {
            "title": "Packaging and quantities",
            "blocks": [
                table(
                    ["", ""],
                    [
                        [cell("Unit"), cell(pk.unit if pk else None)],
                        [cell("Folding"), cell(pk.folding if pk else None)],
                        [cell("Labels"), cell(", ".join(pk.labels) if pk else None)],
                        [cell("Carton"), cell(pk.carton if pk else None)],
                    ],
                    ["1fr", "3fr"],
                ),
                table(
                    ["Size", "Quantity"],
                    [[cell(k), cell(v)] for k, v in p.quantities.items()] or [[GAP, GAP]],
                ),
            ],
        }
    )

    if p.notes:
        s.append({"title": "Notes", "blocks": [para(n) for n in p.notes]})

    rows = [
        [cell(f.severity), cell(f.field), cell(f.message), cell(f.why), cell(f.source)]
        for f in findings
    ]
    s.append(
        {
            "title": "Open gaps",
            "blocks": [
                para(
                    "Everything a factory would still have to ask. Revision-history gaps were "
                    "revised on a past sample: state them now."
                    if findings
                    else "None: this package states everything the checks look for."
                ),
                table(
                    ["Severity", "Field", "Gap", "Why the factory needs it", "Source"],
                    rows,
                    ["auto", "1.3fr", "2fr", "2fr", "auto"],
                )
                if rows
                else para(""),
            ],
        }
    )
    return {"cover": cover, "sections": s}


def render_pdf(
    p: Product, findings: list[Finding], resolve: Resolver | None = None, root: Path | None = None
) -> bytes:
    root = root or Path.cwd()
    v = view(p, findings, resolve, root)
    return typst.compile(
        TEMPLATE.read_bytes(),
        root=str(root),
        sys_inputs={"view": json.dumps(v, sort_keys=True)},
        ignore_system_fonts=True,  # bundled fonts only: the same bytes on every machine
    )


def render_bom(p: Product, findings: list[Finding]) -> bytes:
    """The BOM workbook: materials, trims, hardware, labels, measurements, gaps."""
    wb = Workbook()
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    wb.properties.creator = wb.properties.lastModifiedBy = "design-manufacture-interface"
    gap_fill = PatternFill("solid", fgColor="FFF1B8")
    v = view(p, findings)
    sheets = {sec["title"]: sec for sec in v["sections"]}
    first = True
    for title in ("Bill of materials", "Measurements", "Construction", "Open gaps"):
        sec = sheets.get(title)
        if sec is None:
            continue
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        first = False
        r = 1
        for b in sec["blocks"]:
            if b["kind"] == "para" and b["text"]:
                ws.cell(r, 1, b["text"]).font = Font(bold=True)
                r += 1
            elif b["kind"] == "table":
                for j, col in enumerate(b["columns"], 1):
                    ws.cell(r, j, col).font = Font(bold=True)
                r += 1
                for row in b["rows"]:
                    for j, c in enumerate(row, 1):
                        x = ws.cell(r, j, c["text"])
                        if c["gap"]:
                            x.fill = gap_fill
                    r += 1
                r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return _fixed_zip(buf.getvalue())


def _fixed_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(name))
    return out.getvalue()
